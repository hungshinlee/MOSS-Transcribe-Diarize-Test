"""
Hotword benchmark 評估工具（依賴 openpyxl 跟 cn2an，其餘為 Python 標準庫）。

用法:
    python scripts/evaluate.py --candidate <candidate_dir> \\
        [--predicted-keywords predicted_keywords.json] \\
        [--output report.xlsx]

<candidate_dir> 底下需要有 <audio_id>/transcription.json（ASR 逐字稿），內容格式
跟本專案 pipeline 的輸出一致，支援兩種寫法：
    [{"text": "..."}, ...]   (list of dict，多個 segment 的 text 會被串接)
    {"text": "..."}          (單一 dict)
缺檔的 audio_id 視為候選系統完全沒有輸出（空字串）。
"""
import argparse
import json
import os
import re
import sys
import time
import warnings
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Allow importing from scripts directory directly
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mer_utils import (  # noqa: E402
    build_sibling_superset_map,
    compute_mer,
    count_keyword_occurrences,
    keyword_present,
    levenshtein_distance,
    normalize_cjk_english_spacing,
    normalize_english_digit_words,
    normalize_numbers,
    normalize_spaced_letters,
    preprocess_text,
    tokenize_mixed,
)

HIT_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
MISS_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
EXTRA_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
BOLD = Font(bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# 候選結果讀取
# ---------------------------------------------------------------------------
def read_candidate_text(candidate_dir: Path, audio_id: str) -> str:
    json_path = candidate_dir / audio_id / "transcription.json"
    if not json_path.exists():
        return ""
    try:
        with json_path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            return " ".join(seg.get("text", "") for seg in data)
        if isinstance(data, dict):
            return data.get("text", "")
    except Exception as e:
        print(f"⚠️ 讀取候選結果失敗 {json_path}: {e}")
    return ""


def _safe_output_path(path: Path) -> Path:
    """防呆：若目標檔案已存在，改用帶時間戳記的新檔名，不覆寫舊報表。"""
    if not path.exists():
        return path
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    new_path = path.with_stem(f"{path.stem}_{timestamp}")
    print(f"⚠️ 偵測到 {path} 已存在，為避免覆蓋，本次改輸出到: {new_path}")
    return new_path


def target_keywords_for(keywords: list[str]) -> list[str]:
    """該音檔的目標熱詞（去重、排序）。"""
    return sorted(set(keywords))


def evaluate(candidate_dir: Path, hotwords_map: dict, pseudo_map: dict, true_map: dict, normalize: bool = True):
    audio_ids = sorted(hotwords_map.keys(), key=int)
    sibling_superset_map = build_sibling_superset_map(hotwords_map)

    per_audio = []
    total_tp = total_fn = total_excess = 0
    total_mer_errors = total_mer_ref_len = 0
    total_true_mer_errors = total_true_mer_ref_len = 0

    for aid in audio_ids:
        # 用 Counter 保留每個熱詞在這支音檔裡標記的次數，不去重
        keyword_counts = Counter(hotwords_map.get(aid, []))
        pseudo_text = pseudo_map.get(aid, "")
        hyp_text = read_candidate_text(candidate_dir, aid)

        target_keywords, hits, misses = [], [], []
        for kw in sorted(keyword_counts):
            required = keyword_counts[kw]
            relevant_supersets = [
                s for s in sibling_superset_map.get(kw, []) if s in keyword_counts
            ]
            matched = count_keyword_occurrences(kw, hyp_text, relevant_supersets)
            credited = min(matched, required)  # 講超過標記次數不會多算分
            excess = max(0, matched - required)  # 講超過標記次數的部分，疑似誤報
            total_tp += credited
            total_fn += required - credited
            total_excess += excess
            label = kw if required == matched == 1 else f"{kw} ({matched}/{required})"
            target_keywords.append(kw if required == 1 else f"{kw} x{required}")
            (hits if credited > 0 else misses).append(label)

        # MER 才跟 pseudo 逐字稿比對
        mer_errors, mer_ref_len = compute_mer(pseudo_text, hyp_text, normalize=normalize)
        total_mer_errors += mer_errors
        total_mer_ref_len += mer_ref_len
        mer = mer_errors / mer_ref_len if mer_ref_len else None

        # true_transcribe 跟真人校正過的逐字稿比對
        true_text = true_map.get(aid, "")
        true_mer_errors, true_mer_ref_len = compute_mer(true_text, hyp_text, normalize=normalize)
        total_true_mer_errors += true_mer_errors
        total_true_mer_ref_len += true_mer_ref_len
        true_mer = true_mer_errors / true_mer_ref_len if true_mer_ref_len else None

        per_audio.append(
            {
                "audio_id": aid,
                "target_keywords": target_keywords,
                "hits": hits,
                "misses": misses,
                "mer": mer,
                "pseudo_text": pseudo_text,
                "hyp_text": hyp_text,
                "true_mer": true_mer,
                "true_text": true_text,
            }
        )

    recall = total_tp / (total_tp + total_fn) if (total_tp + total_fn) else 0.0
    overall_mer = (
        total_mer_errors / total_mer_ref_len if total_mer_ref_len else 0.0
    )
    true_transcribe = (
        total_true_mer_errors / total_true_mer_ref_len if total_true_mer_ref_len else None
    )

    summary = {
        "audio_count": len(audio_ids),
        "keyword_instances": total_tp + total_fn,
        "tp": total_tp,
        "fn": total_fn,
        "excess": total_excess,
        "recall": recall,
        "mer": overall_mer,
        "mer_ref_len": total_mer_ref_len,
        "true_transcribe": true_transcribe,
        "true_transcribe_ref_len": total_true_mer_ref_len,
    }
    return summary, per_audio


# ---------------------------------------------------------------------------
# Optional：獨立熱詞預測評估（有明確的預測清單，才有意義的 Precision/F1）
# ---------------------------------------------------------------------------
def _normalize_keyword(kw: str) -> str:
    return kw.strip().casefold()


def evaluate_keyword_predictions(predicted_map: dict, hotwords_map: dict):
    audio_ids = sorted(hotwords_map.keys(), key=int)

    per_audio = []
    total_tp = total_fp = total_fn = 0

    for aid in audio_ids:
        target_keywords = target_keywords_for(hotwords_map.get(aid, []))
        target_norm = {_normalize_keyword(kw): kw for kw in target_keywords}

        predicted_keywords = predicted_map.get(aid, [])
        predicted_norm = {_normalize_keyword(kw): kw for kw in predicted_keywords}

        hits = [orig for norm, orig in target_norm.items() if norm in predicted_norm]
        misses = [orig for norm, orig in target_norm.items() if norm not in predicted_norm]
        extras = [orig for norm, orig in predicted_norm.items() if norm not in target_norm]

        total_tp += len(hits)
        total_fn += len(misses)
        total_fp += len(extras)

        per_audio.append(
            {
                "audio_id": aid,
                "predicted_keywords": sorted(predicted_keywords),
                "hits": sorted(hits),
                "misses": sorted(misses),
                "extras": sorted(extras),
            }
        )

    recall = total_tp / (total_tp + total_fn) if (total_tp + total_fn) else 0.0
    precision = total_tp / (total_tp + total_fp) if (total_tp + total_fp) else 0.0
    f1 = (
        2 * precision * recall / (precision + recall) if (precision + recall) else 0.0
    )

    summary = {
        "tp": total_tp,
        "fp": total_fp,
        "fn": total_fn,
        "recall": recall,
        "precision": precision,
        "f1": f1,
    }
    return summary, per_audio


# ---------------------------------------------------------------------------
# xlsx 報表
# ---------------------------------------------------------------------------
def write_report(
    summary: dict,
    per_audio: list,
    output_path: Path,
    kw_summary: dict | None = None,
    kw_per_audio: list | None = None,
):
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["項目", "數值"])
    for cell in ws_summary[1]:
        cell.font = BOLD
    rows = [
        ("Audio 筆數", summary["audio_count"]),
        ("熱詞出現實例數（Target, 即 TP+FN，來自 hotwords.json）", summary["keyword_instances"]),
        ("[ASR] Hit (TP)", summary["tp"]),
        ("[ASR] Miss (FN)", summary["fn"]),
        (
            "[ASR] 疑似誤報 / 講超過標記次數（matched > required 的超出量加總，"
            "可能是 ASR 多講，也可能是 ground truth 漏標，僅供參考）",
            summary["excess"],
        ),
        ("[ASR] Recall（子字串比對候選逐字稿）", f"{summary['recall']:.2%}"),
        ("[ASR] MER（跟 pseudo 逐字稿比對，中文字/英文單字/數字各當一個 token 算編輯距離，加減參考）", f"{summary['mer']:.2%}"),
        ("[ASR] MER 分母（pseudo token 數）", summary["mer_ref_len"]),
        (
            f"[ASR] True Transcribe MER（跟真人校正逐字稿比較，"
            f"{summary['true_transcribe_ref_len']} tokens，僅涵蓋有標註真人逐字稿的音檔）",
            f"{summary['true_transcribe']:.2%}" if summary["true_transcribe"] is not None else "N/A",
        ),
    ]
    for label, value in rows:
        ws_summary.append([label, value])

    if kw_summary is not None:
        ws_summary.append([])
        ws_summary.append(["熱詞預測評估 (optional, --predicted-keywords)", ""])
        ws_summary.cell(row=ws_summary.max_row, column=1).font = BOLD
        kw_rows = [
            ("[KW] Hit (TP)", kw_summary["tp"]),
            ("[KW] Miss (FN)", kw_summary["fn"]),
            ("[KW] Extra / 誤預測 (FP)", kw_summary["fp"]),
            ("[KW] Recall", f"{kw_summary['recall']:.2%}"),
            ("[KW] Precision", f"{kw_summary['precision']:.2%}"),
            ("[KW] F1-Score", round(kw_summary["f1"], 4)),
        ]
        for label, value in kw_rows:
            ws_summary.append([label, value])
    else:
        ws_summary.append([])
        ws_summary.append(
            [
                "（未提供 --predicted-keywords，略過熱詞預測評估。ASR Recall 不附 "
                "Precision/F1：自由文字裡沒有可靠方法判定候選「宣稱」了哪些熱詞。）"
            ]
        )
        ws_summary.cell(row=ws_summary.max_row, column=1).alignment = Alignment(wrap_text=True)
        ws_summary.merge_cells(
            start_row=ws_summary.max_row, start_column=1, end_row=ws_summary.max_row, end_column=2
        )

    ws_summary.column_dimensions["A"].width = 45
    ws_summary.column_dimensions["B"].width = 20

    ws_detail = wb.create_sheet("Per-Audio")
    header = [
        "Audio ID",
        "Target Keywords",
        "[ASR] Hit Keywords",
        "[ASR] Missed Keywords",
        "MER (vs pseudo)",
        "Pseudo Reference",
        "True Transcribe MER",
        "True Reference",
        "Candidate Text",
    ]
    has_kw = kw_per_audio is not None
    if has_kw:
        header += [
            "Predicted Keywords",
            "[KW] Hit (TP)",
            "[KW] Miss (FN)",
            "[KW] Extra (FP)",
        ]
        kw_by_id = {rec["audio_id"]: rec for rec in kw_per_audio}

    ws_detail.append(header)
    for cell in ws_detail[1]:
        cell.font = BOLD

    text_wrap_cols = [2, 3, 4, 6, 8, 9] + ([10, 11, 12, 13] if has_kw else [])
    for rec in per_audio:
        mer_display = f"{rec['mer']:.2%}" if rec["mer"] is not None else "N/A"
        true_mer_display = f"{rec['true_mer']:.2%}" if rec["true_mer"] is not None else "N/A"
        row = [
            int(rec["audio_id"]),
            ", ".join(rec["target_keywords"]),
            ", ".join(rec["hits"]),
            ", ".join(rec["misses"]),
            mer_display,
            rec["pseudo_text"],
            true_mer_display,
            rec["true_text"],
            rec["hyp_text"],
        ]
        if has_kw:
            kw_rec = kw_by_id.get(rec["audio_id"], {})
            row += [
                ", ".join(kw_rec.get("predicted_keywords", [])),
                ", ".join(kw_rec.get("hits", [])),
                ", ".join(kw_rec.get("misses", [])),
                ", ".join(kw_rec.get("extras", [])),
            ]

        ws_detail.append(row)
        r = ws_detail.max_row
        if rec["hits"]:
            ws_detail.cell(row=r, column=3).fill = HIT_FILL
        if rec["misses"]:
            ws_detail.cell(row=r, column=4).fill = MISS_FILL
        if has_kw:
            kw_rec = kw_by_id.get(rec["audio_id"], {})
            if kw_rec.get("hits"):
                ws_detail.cell(row=r, column=11).fill = HIT_FILL
            if kw_rec.get("misses"):
                ws_detail.cell(row=r, column=12).fill = MISS_FILL
            if kw_rec.get("extras"):
                ws_detail.cell(row=r, column=13).fill = EXTRA_FILL
        for col in text_wrap_cols:
            ws_detail.cell(row=r, column=col).alignment = WRAP_TOP

    widths = {1: 10, 2: 26, 3: 22, 4: 22, 5: 12, 6: 60, 7: 14, 8: 60, 9: 60, 10: 26, 11: 18, 12: 18, 13: 18}
    for col, width in widths.items():
        if col > len(header):
            continue
        ws_detail.column_dimensions[ws_detail.cell(row=1, column=col).column_letter].width = width
    ws_detail.freeze_panes = "A2"

    output_path = _safe_output_path(output_path)
    wb.save(output_path)
    print(f"📊 報表已產生: {output_path.resolve()}")


def main():
    parser = argparse.ArgumentParser(description="Hotword benchmark 評估工具")
    parser.add_argument(
        "--candidate",
        required=True,
        help="候選系統結果資料夾，底下需有 <audio_id>/transcription.json",
    )
    parser.add_argument(
        "--benchmark-dir",
        default=None,
        help="hotword_benchmark 資料夾路徑，包含 hotwords.json, pseudo_transcripts.json 等",
    )
    parser.add_argument(
        "--predicted-keywords",
        default=None,
        help=(
            "optional: 候選熱詞預測 JSON 檔路徑，格式為 "
            '{"<audio_id>": ["熱詞1", "熱詞2", ...], ...}。'
            "提供時會額外算出真正獨立的 Precision/Recall/F1。"
        ),
    )
    parser.add_argument(
        "--output", default="report.xlsx", help="輸出報表路徑 (預設 report.xlsx)"
    )
    parser.add_argument(
        "--no-normalize",
        action="store_true",
        help=(
            "關閉 MER 計算的文字正規化（合併空格拼字母、中文數字轉阿拉伯數字），"
            "預設是開啟的"
        ),
    )
    args = parser.parse_args()

    candidate_dir = Path(args.candidate).resolve()
    if not candidate_dir.exists():
        raise FileNotFoundError(f"找不到候選結果資料夾: {candidate_dir}")

    here = Path(__file__).resolve().parent
    repro_root = here.parent
    if args.benchmark_dir:
        bench_dir = Path(args.benchmark_dir).resolve()
    elif (here / "hotwords.json").exists():
        bench_dir = here
    else:
        bench_dir = repro_root / "data" / "hotword_benchmark"

    hotwords_file = bench_dir / "hotwords.json"
    pseudo_file = bench_dir / "pseudo_transcripts.json"
    true_file = bench_dir / "true_transcripts" / "true_transcripts.json"

    if not hotwords_file.exists():
        raise FileNotFoundError(f"找不到基準檔案: {hotwords_file}")

    with hotwords_file.open("r", encoding="utf-8") as f:
        hotwords_map = json.load(f)
    pseudo_map = json.load(pseudo_file.open("r", encoding="utf-8")) if pseudo_file.exists() else {}
    true_map = json.load(true_file.open("r", encoding="utf-8")) if true_file.exists() else {}

    summary, per_audio = evaluate(
        candidate_dir, hotwords_map, pseudo_map, true_map, normalize=not args.no_normalize
    )

    print("=" * 60)
    print(f"Audio 筆數       : {summary['audio_count']}")
    print(f"[ASR] Recall (vs hotwords.json): {summary['recall']:.2%}")
    print(f"[ASR] MER (vs pseudo, 加減參考): {summary['mer']:.2%}")
    print(
        f"[ASR] True Transcribe MER (vs 真人逐字稿, {summary['true_transcribe_ref_len']} tokens): "
        + (f"{summary['true_transcribe']:.2%}" if summary["true_transcribe"] is not None else "N/A")
    )

    kw_summary = kw_per_audio = None
    if args.predicted_keywords:
        predicted_path = Path(args.predicted_keywords).resolve()
        if not predicted_path.exists():
            raise FileNotFoundError(f"找不到熱詞預測檔案: {predicted_path}")
        with predicted_path.open("r", encoding="utf-8") as f:
            predicted_map = json.load(f)
        kw_summary, kw_per_audio = evaluate_keyword_predictions(predicted_map, hotwords_map)
        print(f"[KW]  Recall     : {kw_summary['recall']:.2%}")
        print(f"[KW]  Precision  : {kw_summary['precision']:.2%}")
        print(f"[KW]  F1-Score   : {kw_summary['f1']:.4f}")
    print("=" * 60)

    write_report(summary, per_audio, Path(args.output), kw_summary, kw_per_audio)


if __name__ == "__main__":
    main()
